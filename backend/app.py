import io

import hashlib

import json

import os

import re

import sqlite3

import threading

import uuid

from pathlib import Path

from contextlib import contextmanager

from fastapi import FastAPI, File, Form, HTTPException, UploadFile

from fastapi.middleware.cors import CORSMiddleware

from fastapi.responses import FileResponse

from pydantic import BaseModel, Field

from pypdf import PdfReader

from docx import Document

BASE = Path(__file__).resolve().parent

FRONTEND = BASE.parent / "frontend" / "index.html"

DB = BASE / "legal_chat.db"

DATA = BASE / "data"

LIBRARY_SESSION = "__library__"
INDEX_VERSION = "v4-numbered-headings"

MAX_FILE = 10 * 1024 * 1024

MAX_SESSION = 100 * 1024 * 1024

MAX_TOTAL = 1024 * 1024 * 1024

MAX_PDF_PAGES = 1500

DISCLAIMER = (

    "For legal education and research only. "

    "Verify legal information against current official sources. "

    "Consult a qualified lawyer for advice about your situation."

)

app = FastAPI(

    title="Legal AI Assistant",

    description="Indian legal document research and study",

    version="3.0"

)

app.add_middleware(

    CORSMiddleware,

    allow_origins=[

        "http://127.0.0.1:5500",

        "http://localhost:5500",

        "http://127.0.0.1:8000",

        "http://localhost:8000"

    ],

    allow_methods=["*"],

    allow_headers=["*"]

)

LOCK = threading.RLock()

@contextmanager

def database():

    with LOCK:

        conn = sqlite3.connect(DB, timeout=30)

        conn.row_factory = sqlite3.Row

        try:

            yield conn

            conn.commit()

        finally:

            conn.close()

def initialize_database():

    DATA.mkdir(exist_ok=True)

    with database() as conn:

        conn.execute("""

            CREATE TABLE IF NOT EXISTS documents (

                id TEXT PRIMARY KEY,

                session TEXT NOT NULL,

                filename TEXT NOT NULL,

                size INTEGER NOT NULL,

                created_at TEXT DEFAULT CURRENT_TIMESTAMP

            )

        """)

        conn.execute("""

            CREATE TABLE IF NOT EXISTS passages (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                document_id TEXT NOT NULL,

                page INTEGER,

                content TEXT NOT NULL

            )

        """)

        conn.execute("""

            CREATE TABLE IF NOT EXISTS messages (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                session TEXT NOT NULL,

                role TEXT NOT NULL,

                content TEXT NOT NULL,

                citations TEXT DEFAULT '[]',

                created_at TEXT DEFAULT CURRENT_TIMESTAMP

            )

        """)

        # Fix older SQLite databases without these columns.

        migrations = {

            "messages": {

                "citations": "TEXT DEFAULT '[]'",

                "created_at": "TEXT"

            },

            "documents": {

                "created_at": "TEXT",

                "stored_path": "TEXT",

                "content_hash": "TEXT"

            }

        }

        for table, columns in migrations.items():

            existing = {

                row[1]

                for row in conn.execute(

                    f"PRAGMA table_info({table})"

                )

            }

            for name, definition in columns.items():

                if name not in existing:

                    conn.execute(

                        f"ALTER TABLE {table} "

                        f"ADD COLUMN {name} {definition}"

                    )

        conn.execute("""

            CREATE INDEX IF NOT EXISTS idx_doc_session

            ON documents(session)

        """)

        conn.execute("""

            CREATE INDEX IF NOT EXISTS idx_passages_document

            ON passages(document_id)

        """)

        conn.execute("""

            CREATE INDEX IF NOT EXISTS idx_messages_session

            ON messages(session)

        """)

initialize_database()

class Question(BaseModel):

    query: str = Field(min_length=1, max_length=3000)

    session_id: str = Field(

        default="local",

        pattern=r"^[A-Za-z0-9_-]{1,120}$"

    )

def validate_session(session_id):

    if not re.fullmatch(

        r"[A-Za-z0-9_-]{1,120}",

        session_id

    ):

        raise HTTPException(400, "Invalid session ID")

    return session_id

def clean_text(text):

    return re.sub(r"\s+", " ", text).strip()

def split_text(text, size=1400, overlap=220):

    text = clean_text(text)

    if not text:

        return []

    result = []

    start = 0

    while start < len(text):

        end = min(start + size, len(text))

        if end < len(text):

            boundary = text.rfind(

                " ", start + 400, end

            )

            if boundary > start:

                end = boundary

        part = text[start:end].strip()

        if part:

            result.append(part)

        if end >= len(text):

            break

        start = max(start + 1, end - overlap)

    return result

def extract_document(filename, raw):

    extension = Path(filename).suffix.lower()

    try:

        if extension == ".pdf":

            reader = PdfReader(io.BytesIO(raw))

            if reader.is_encrypted:

                raise HTTPException(

                    400, "Encrypted PDFs are not supported."

                )

            if len(reader.pages) > MAX_PDF_PAGES:

                raise HTTPException(

                    400,

                    f"PDF exceeds {MAX_PDF_PAGES} pages."

                )

            return [

                (i + 1, page.extract_text() or "")

                for i, page in enumerate(reader.pages)

            ]

        if extension == ".docx":

            doc = Document(io.BytesIO(raw))

            paragraphs = [

                p.text for p in doc.paragraphs

            ]

            for table in doc.tables:

                for row in table.rows:

                    paragraphs.append(

                        " | ".join(

                            cell.text for cell in row.cells

                        )

                    )

            return [(None, "\n".join(paragraphs))]

        if extension == ".txt":

            return [

                (None, raw.decode("utf-8-sig"))

            ]

    except HTTPException:

        raise

    except Exception:

        raise HTTPException(

            400,

            "Cannot read document. Upload a valid "

            "PDF, DOCX or UTF-8 TXT file."

        )

    raise HTTPException(

        400, "Only PDF, DOCX and TXT are supported."

    )

STOP_WORDS = set("""

a an the and or is are was were to for from with

what which where when why how does did do can could

would should please tell give explain describe me my

of in on at by as it this that about under

india indian law legal constitution article section

provision provisions meaning definition

""".split())

def keywords(text):

    words = re.findall(

        r"[a-z0-9]+", text.lower()

    )

    return [

        word for word in words

        if word not in STOP_WORDS

        and (len(word) > 2 or word.isdigit())

    ]

LEGAL_PATTERN = re.compile(

    r"\b(law|legal|constitution|article|section|"

    r"court|judge|judgment|case|rights?|crime|"

    r"criminal|civil|contract|police|bail|"

    r"petition|preamble|fundamental|property|"

    r"divorce|consumer|employment|tax|"

    r"parliament|amendment|ipc|bns|bnss|"

    r"evidence|statute|act|policy|citizen|"

    r"justice|procedure|directive)\b",

    re.I

)

NONLEGAL_PATTERN = re.compile(

    r"\b(recipe|cook|weather|football score|"

    r"movie recommendation|write a poem|"

    r"coding tutorial)\b",

    re.I

)

UNSAFE_PATTERN = re.compile(

    r"\b(how to|steps to|help me|ways to|"

    r"teach me|instructions to)\b"

    r".{0,100}\b("

    r"forge|bribe|fabricate evidence|"

    r"hide evidence|fake documents|"

    r"intimidate witness|launder money|"

    r"evade police)\b",

    re.I

)

def extract_references(query):

    articles = re.findall(

        r"\barticle\s+(\d+[a-z]?)\b",

        query, re.I

    )

    sections = re.findall(

        r"\bsection\s+(\d+[a-z]?)\b",

        query, re.I

    )

    return articles, sections

def reference_match(content, kind, number):
    """Match references, including bare numbered headings in Indian legal PDFs."""
    number = re.escape(number)
    prefix = r"(?:article|art\.?)" if kind == "article" else r"(?:section|sec\.?)"
    explicit = rf"\b{prefix}\s*{number}\b"
    heading = rf"(?:^|(?<!\d)\s){number}\s*[.\-–:]\s*[A-Z]"
    return bool(re.search(explicit, content, re.I) or re.search(heading, content))


def relevance_score(query, passage):
    content = passage.lower()
    query_words = set(keywords(query))
    passage_words = set(keywords(content))
    score = len(query_words & passage_words) * 3
    articles, sections = extract_references(query)
    for number in articles:
        if reference_match(passage, "article", number):
            score += 60
        else:
            score -= 40
    for number in sections:
        if reference_match(passage, "section", number):
            score += 60
        else:
            score -= 40
    query_clean = clean_text(query).lower()
    if len(query_clean) > 10 and query_clean in content:
        score += 15
    return score


def retrieve(conn, session, query, limit=5):

    rows = conn.execute("""

        SELECT

            p.page,

            p.content,

            d.filename,

            d.id

        FROM passages p

        JOIN documents d

            ON d.id = p.document_id

        WHERE d.session IN (?, ?)

    """, (session, LIBRARY_SESSION)).fetchall()

    ranked = []

    for row in rows:

        score = relevance_score(

            query, row["content"]

        )

        if score > 0:

            ranked.append((score, dict(row)))

    articles, sections = extract_references(query)
    if articles or sections:
        matching = [(score, row) for score, row in ranked
                    if all(reference_match(row["content"], "article", n) for n in articles)
                    and all(reference_match(row["content"], "section", n) for n in sections)]
        if matching:
            ranked = matching
    ranked.sort(

        key=lambda item: item[0],

        reverse=True

    )

    seen = set()

    result = []

    for score, row in ranked:

        fingerprint = clean_text(row["content"]).lower()[:250]

        if fingerprint in seen:

            continue

        seen.add(fingerprint)

        result.append(row)

        if len(result) >= limit:

            break

    return result

def save_message(conn, session, role, content, citations):

    conn.execute("""

        INSERT INTO messages

        (session, role, content, citations)

        VALUES (?, ?, ?, ?)

    """, (

        session,

        role,

        content,

        json.dumps(citations)

    ))

def sync_data_folder():

    """Keep shared library in sync with backend/data (add/change/remove)."""

    DATA.mkdir(parents=True, exist_ok=True)

    with database() as conn:

        current = {p.name: p for p in DATA.iterdir()

                   if p.is_file() and p.suffix.lower() in {".pdf", ".docx", ".txt"}}

        # Clean up legacy duplicate rows and stale records, without touching chat history.

        rows = conn.execute("SELECT id, session, filename, stored_path, content_hash FROM documents").fetchall()

        for row in rows:

            name = row["stored_path"]

            if row["session"] == LIBRARY_SESSION and name and name not in current:

                conn.execute("DELETE FROM passages WHERE document_id=?", (row["id"],))

                conn.execute("DELETE FROM documents WHERE id=?", (row["id"],))

        for name, path in sorted(current.items()):

            try:

                # Legacy versions could index the same file without stored_path.

                old_duplicates = conn.execute(

                    "SELECT id FROM documents WHERE filename=? AND stored_path IS NULL",

                    (name,)).fetchall()

                for duplicate in old_duplicates:

                    conn.execute("DELETE FROM passages WHERE document_id=?", (duplicate["id"],))

                    conn.execute("DELETE FROM documents WHERE id=?", (duplicate["id"],))

                if path.stat().st_size > MAX_FILE or path.stat().st_size == 0:

                    continue

                raw = path.read_bytes()

                digest = hashlib.sha256(raw + INDEX_VERSION.encode()).hexdigest()

                rows = conn.execute(

                    "SELECT id, content_hash FROM documents WHERE stored_path=? ORDER BY rowid",

                    (name,)).fetchall()

                # Keep one shared document per physical file. Existing session-upload

                # entries are migrated into the shared library rather than duplicated.

                if rows:

                    primary = rows[0]["id"]

                    for duplicate in rows[1:]:

                        conn.execute("DELETE FROM passages WHERE document_id=?", (duplicate["id"],))

                        conn.execute("DELETE FROM documents WHERE id=?", (duplicate["id"],))

                    if rows[0]["content_hash"] == digest:

                        conn.execute("UPDATE documents SET session=?, filename=?, size=? WHERE id=?",

                                     (LIBRARY_SESSION, name, len(raw), primary))

                        continue

                else:

                    primary = uuid.uuid4().hex

                pages = extract_document(name, raw)

                chunks = [(page, chunk) for page, content in pages for chunk in split_text(content)]

                if not chunks:

                    # No searchable text: leave the previous index intact if there is one.

                    continue

                if rows:

                    conn.execute("DELETE FROM passages WHERE document_id=?", (primary,))

                    conn.execute("UPDATE documents SET session=?, filename=?, size=?, content_hash=? WHERE id=?",

                                 (LIBRARY_SESSION, name, len(raw), digest, primary))

                else:

                    conn.execute("INSERT INTO documents (id, session, filename, size, stored_path, content_hash) "

                                 "VALUES (?, ?, ?, ?, ?, ?)",

                                 (primary, LIBRARY_SESSION, name, len(raw), name, digest))

                conn.executemany("INSERT INTO passages (document_id, page, content) VALUES (?, ?, ?)",

                                 [(primary, page, chunk) for page, chunk in chunks])

            except (HTTPException, OSError, ValueError):

                continue

sync_data_folder()

@app.get("/")

def root():

    return {

        "status": "Legal AI API running",

        "version": "3.0"

    }

@app.get("/health")

def health():

    return {"status": "online"}

@app.get("/app")

def frontend():

    if not FRONTEND.exists():

        raise HTTPException(

            404, "frontend/index.html not found."

        )

    return FileResponse(FRONTEND)

@app.get("/documents")

def list_documents(session_id: str = "local"):

    validate_session(session_id)

    with database() as conn:

        documents = conn.execute("""

            SELECT id, filename, size

            FROM documents

            WHERE session IN (?, ?)

            ORDER BY created_at DESC

        """, (session_id, LIBRARY_SESSION)).fetchall()

        session_bytes = conn.execute("""

            SELECT COALESCE(SUM(size), 0)

            FROM documents

            WHERE session = ?

        """, (session_id,)).fetchone()[0]

        total_bytes = conn.execute("""

            SELECT COALESCE(SUM(size), 0)

            FROM documents

        """).fetchone()[0]

    return {

        "documents": [

            dict(row) for row in documents

        ],

        "session_bytes": session_bytes,

        "total_bytes": total_bytes,

        "max_file_bytes": MAX_FILE,

        "max_session_bytes": MAX_SESSION,

        "max_total_bytes": MAX_TOTAL

    }

@app.post("/upload")

async def upload_document(

    session_id: str = Form(...),

    file: UploadFile = File(...)

):

    validate_session(session_id)

    filename = Path(

        (file.filename or "").replace("\\\\\\\\\\\\\\\\", "/")

    ).name[:180]

    if Path(filename).suffix.lower() not in {

        ".pdf", ".docx", ".txt"

    }:

        raise HTTPException(

            400, "Unsupported file type."

        )

    raw = await file.read(MAX_FILE + 1)

    await file.close()

    if not raw or len(raw) > MAX_FILE:

        raise HTTPException(

            413, "Maximum file size is 10 MB."

        )

    pages = extract_document(filename, raw)

    passages = [

        (page, chunk)

        for page, text in pages

        for chunk in split_text(text)

    ]

    if not passages:

        raise HTTPException(

            400,

            "No selectable text found. "

            "Scanned PDFs need OCR."

        )

    # Uploading through the UI is optional: files copied into data/ work too.

    # Save under the original filename, without overwriting existing documents.

    safe_name = re.sub(r"[^A-Za-z0-9._ -]", "_", filename).strip(" .")

    if not safe_name:

        raise HTTPException(400, "Invalid filename")

    target = DATA / safe_name

    with database() as conn:

        total_used = conn.execute("SELECT COALESCE(SUM(size),0) FROM documents").fetchone()[0]

    if total_used + len(raw) > MAX_TOTAL:

        raise HTTPException(413, "Total storage limit exceeded")

    if target.exists():

        if hashlib.sha256(target.read_bytes()).digest() == hashlib.sha256(raw).digest():

            sync_data_folder()

            return {"filename": safe_name, "size": len(raw), "chunks": len(passages), "already_exists": True}

        stem, suffix = Path(safe_name).stem, Path(safe_name).suffix

        counter = 2

        while target.exists():

            target = DATA / f"{stem}_{counter}{suffix}"

            counter += 1

    target.write_bytes(raw)

    sync_data_folder()

    return {"filename": target.name, "size": len(raw), "chunks": len(passages)}

@app.delete("/documents/{document_id}")

def delete_document(

    document_id: str,

    session_id: str

):

    validate_session(session_id)

    with database() as conn:

        found = conn.execute("""

            SELECT id, stored_path FROM documents

            WHERE id = ? AND session IN (?, ?)

        """, (

            document_id, session_id, LIBRARY_SESSION

        )).fetchone()

        if not found:

            raise HTTPException(

                404, "Document not found."

            )

        conn.execute("""

            DELETE FROM passages

            WHERE document_id = ?

        """, (document_id,))

        conn.execute("""

            DELETE FROM documents

            WHERE id = ? AND session IN (?, ?)

        """, (

            document_id, session_id, LIBRARY_SESSION

        ))

    if found["stored_path"]:

        (DATA / Path(found["stored_path"]).name).unlink(missing_ok=True)

    return {"deleted": True}

@app.get("/history")

def get_history(session_id: str = "local"):

    validate_session(session_id)

    with database() as conn:

        rows = conn.execute("""

            SELECT role, content, citations

            FROM messages

            WHERE session = ?

            ORDER BY id DESC

            LIMIT 100

        """, (session_id,)).fetchall()

    messages = []

    for row in reversed(rows):

        try:

            citations = json.loads(

                row["citations"] or "[]"

            )

        except (ValueError, TypeError):

            citations = []

        messages.append({

            "role": row["role"],

            "content": row["content"],

            "citations": citations

        })

    return {"messages": messages}

@app.delete("/history")

def clear_history(session_id: str = "local"):

    validate_session(session_id)

    with database() as conn:

        conn.execute("""

            DELETE FROM messages

            WHERE session = ?

        """, (session_id,))

    return {"deleted": True}

@app.post("/ask")

def ask(question: Question):

    sync_data_folder()

    query = question.query.strip()

    session = question.session_id

    if not query:

        raise HTTPException(

            400, "Question cannot be empty."

        )

    with database() as conn:

        if UNSAFE_PATTERN.search(query):

            answer = (

                "I cannot provide instructions for "

                "illegal activities or falsifying evidence. "

                "I can explain the relevant legal rules, "

                "penalties, rights and lawful procedures."

            )

            citations = []

        elif (

            NONLEGAL_PATTERN.search(query)

            and not LEGAL_PATTERN.search(query)

        ):

            answer = (

                "I specialize in legal education "

                "and research. Please ask a question "

                "about law or the available legal documents."

            )

            citations = []

        else:

            found = retrieve(

                conn, session, query

            )

            if not found:

                answer = (

                    "I could not find sufficiently relevant "

                    "information in the available legal "

                    "documents. Add a relevant file to backend/data/ "

                    "legal text or ask a more specific question."

                )

                citations = []

            else:

                citations = []

                excerpts = []

                for index, row in enumerate(found, 1):

                    citations.append({

                        "number": index,

                        "filename": row["filename"],

                        "page": row["page"],

                        "document_id": row["id"]

                    })

                    excerpt = row["content"][:900]
                    articles, sections = extract_references(query)
                    for kind, numbers in (("article", articles), ("section", sections)):
                        for number in numbers:
                            pattern = rf"\b(?:{kind}|{'art' if kind == 'article' else 'sec'}\.?)\s*{re.escape(number)}\b|(?<!\d)\b{re.escape(number)}\s*[.\-–:]\s*[A-Z]"
                            match = re.search(pattern, row["content"], re.I)
                            if match:
                                excerpt = row["content"][max(0, match.start()-60):match.start()+1000]
                                break

                    excerpts.append(

                        f"[{index}] {excerpt}"

                    )

                answer = (

                    "Relevant legal document passages:\n\n"

                    + "\n\n".join(excerpts)

                    + "\n\nThese are source extracts, "

                    "not independently verified legal "

                    "conclusions. Check the cited provisions "

                    "against the current official law."

                )

        save_message(

            conn, session, "user", query, []

        )

        save_message(

            conn, session, "assistant",

            answer, citations

        )

    return {

        "question": query,

        "answer": answer,

        "summary": answer,

        "citations": citations,

        "disclaimer": DISCLAIMER

    }
